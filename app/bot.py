"""Telegram handlers: menu, legal chat, document analysis, paid document drafting, receipt payments."""
import hashlib
import os
import html
import io
import json
import logging

from aiogram import Bot, F, Router
from aiogram.filters import Command, CommandObject, CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (BufferedInputFile, CallbackQuery, InlineKeyboardButton, InlineKeyboardMarkup,
                           KeyboardButton, Message, ReplyKeyboardMarkup)
from aiogram.utils.chat_action import ChatActionSender
from docx import Document as Docx
from google.genai import errors as genai_errors

from app import config, db, llm, payments, rag

log = logging.getLogger(__name__)
router = Router()

MAX_FILE = 20 * 1024 * 1024  # Telegram Bot API download limit
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FILE_MIMES = {"application/pdf", "image/jpeg", "image/png", "image/webp", "image/heic", DOCX_MIME}


class St(StatesGroup):
    receipt = State()   # waiting for a payment receipt screenshot
    doc_desc = State()  # waiting for "what document do you need"


DISCLAIMER = ("⚠️ Men sun'iy intellektman, litsenziyali advokat emasman. Javoblarim lex.uz'dagi rasmiy qonun "
              "matnlariga asoslanadi va manba havolasi bilan beriladi, lekin muhim qarorlardan oldin advokat "
              "bilan maslahatlashing.")

SITUATIONS = {
    "👷 Mehnat": "Mehnat munosabatlari bo'yicha yordam beraman. Iltimos, yozing:\n• Nima bo'ldi (ishdan bo'shatish, "
                "ish haqi to'lanmagan, ta'til, shartnoma...)?\n• Rasmiy mehnat shartnomasi bormi?\n• Qachon sodir bo'ldi?",
    "🏠 Uy-joy, ijara": "Uy-joy va ijara bo'yicha yordam beraman. Yozing:\n• Siz ijarachimisiz yoki uy egasi?\n"
                       "• Yozma shartnoma bormi?\n• Muammo nimada (chiqarib yuborish, depozit, to'lov, ta'mir...)?",
    "👨‍👩‍👧 Oila, ajrashish": "Oila masalalari bo'yicha yordam beraman. Yozing:\n• Masala nima (ajrashish, aliment, "
                            "bolalar, mol-mulk bo'lish)?\n• Nikoh FHDYOda qayd etilganmi?\n• Voyaga yetmagan bolalar bormi?",
    "🚗 Jarima, YHXX": "Jarima va yo'l harakati bo'yicha yordam beraman. Yozing:\n• Qanday jarima yoki qoidabuzarlik?\n"
                      "• Qachon va qanday rasmiylashtirildi (kamera, inspektor)?\n• Rozimisiz yoki shikoyat qilmoqchimisiz?",
    "🛒 Iste'molchi": "Iste'molchi huquqlari bo'yicha yordam beraman. Yozing:\n• Nima sotib oldingiz / qanday xizmat?\n"
                     "• Muammo nimada (sifatsiz, qaytarib olishmayapti, kafolat)?\n• Chek yoki shartnoma bormi?",
    "📜 Meros": "Meros masalalari bo'yicha yordam beraman. Yozing:\n• Meros qoldiruvchi kim va qachon vafot etgan?\n"
               "• Vasiyatnoma bormi?\n• Merosxo'rlar kimlar va nizo nimada?",
}

MENU = ReplyKeyboardMarkup(resize_keyboard=True, keyboard=[
    [KeyboardButton(text="💬 Savol berish"), KeyboardButton(text="📄 Hujjatni tekshirish")],
    [KeyboardButton(text="👷 Mehnat"), KeyboardButton(text="🏠 Uy-joy, ijara")],
    [KeyboardButton(text="👨‍👩‍👧 Oila, ajrashish"), KeyboardButton(text="🚗 Jarima, YHXX")],
    [KeyboardButton(text="🛒 Iste'molchi"), KeyboardButton(text="📜 Meros")],
    [KeyboardButton(text="📝 Hujjat tayyorlatish"), KeyboardButton(text="💳 Balans")],
    [KeyboardButton(text="🆕 Yangi suhbat")],
])


def _price() -> str:
    return f"{config.DOC_PRICE:,}".replace(",", " ")


def offer_kb() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(
        text=f"📝 Hujjatni tayyorlab berish — {_price()} so'm", callback_data="draft")]])


def split_text(text: str, size: int = 4000) -> list[str]:
    """Telegram messages are limited to 4096 chars; split on line boundaries."""
    parts, buf = [], ""
    for line in text.split("\n"):
        while len(line) > size:
            if buf:
                parts.append(buf)
                buf = ""
            parts.append(line[:size])
            line = line[size:]
        if buf and len(buf) + len(line) + 1 > size:
            parts.append(buf)
            buf = ""
        buf = f"{buf}\n{line}" if buf else line
    if buf.strip():
        parts.append(buf)
    return parts


async def send_long(message: Message, text: str, parse_mode: str | None = None, **kwargs) -> None:
    """kwargs (keyboards) go on the last part only. HTML tags never span lines, so line splits keep them valid."""
    parts = split_text(text)
    for i, part in enumerate(parts):
        await message.answer(part, disable_web_page_preview=True, parse_mode=parse_mode,
                             **(kwargs if i == len(parts) - 1 else {}))


async def ai_error(message: Message, exc: Exception) -> None:
    if isinstance(exc, genai_errors.APIError) and exc.code == 429:
        await message.answer("⏳ Hozir so'rovlar juda ko'p. 1-2 daqiqadan keyin qayta urinib ko'ring.")
    else:
        log.exception("AI request failed")
        await message.answer("❌ Xatolik yuz berdi. Birozdan keyin qayta urinib ko'ring.")


async def quota_ok(message: Message) -> bool:
    if await db.take_daily_quota(message.from_user.id, config.FREE_DAILY_LIMIT):
        return True
    await message.answer(f"Bugungi bepul limit ({config.FREE_DAILY_LIMIT} ta so'rov) tugadi. Ertaga yana yozing 🙏")
    return False


async def reply_with_offer(message: Message, reply: rag.Reply) -> None:
    if reply.offer:
        await db.set_pending_doc(message.from_user.id, reply.offer)
        await send_long(message, reply.text + f"\n\n📝 <b>Kerakli hujjat:</b> {html.escape(reply.offer)}\n"
                                              "Tayyor Word faylini yozib beraman 👇",
                        parse_mode="HTML", reply_markup=offer_kb())
    else:
        await send_long(message, reply.text, parse_mode="HTML")


# ---------------- menu ----------------

@router.message(CommandStart())
async def start(message: Message, state: FSMContext) -> None:
    await state.clear()
    await db.ensure_user(message.from_user.id)
    await message.answer(
        "Assalomu alaykum! Men O'zbekiston qonunchiligi bo'yicha AI huquqiy yordamchiman.\n\n"
        "• Vaziyatingizni yozing — qonun bo'yicha nima qilish kerakligini moddalar bilan tushuntiraman.\n"
        "• Shartnoma yoki hujjat rasmini / PDF'ini yuboring — tahlil qilib, xavfli joylarini ko'rsataman.\n"
        "• Ariza, shikoyat, da'vo arizasini tayyorlab beraman.\n\n" + DISCLAIMER,
        reply_markup=MENU,
    )


@router.message(F.text == "🆕 Yangi suhbat")
@router.message(Command("new"))
async def new_chat(message: Message, state: FSMContext) -> None:
    await state.clear()
    await db.ensure_user(message.from_user.id)
    await db.clear_history(message.from_user.id)
    await db.set_pending_doc(message.from_user.id, None)
    await message.answer("Yangi suhbat boshlandi. Savolingizni yozing.", reply_markup=MENU)


@router.message(F.text == "💬 Savol berish")
async def ask(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer("Vaziyatingizni batafsil yozing: nima bo'ldi, qachon, kim bilan, nimani xohlaysiz.")


@router.message(F.text == "📄 Hujjatni tekshirish")
async def check_doc(message: Message, state: FSMContext) -> None:
    await state.clear()
    await message.answer("Hujjatni yuboring: rasm (har bir sahifa alohida), PDF yoki Word (.docx) fayl.\n"
                         "Rasm/fayl izohiga savolingizni ham yozishingiz mumkin, masalan: "
                         "«Shu ijara shartnomasini imzolasam bo'ladimi?»")


@router.message(F.text.in_(SITUATIONS))
async def situation(message: Message, state: FSMContext) -> None:
    await state.clear()
    prompt = SITUATIONS[message.text]
    await db.add_message(message.from_user.id, "model", prompt)  # so the next answer knows the topic
    await message.answer(prompt)


@router.message(F.text == "💳 Balans")
async def balance(message: Message) -> None:
    user = await db.ensure_user(message.from_user.id)
    await message.answer(
        f"💳 Hujjat tayyorlash kreditlari: {user['credits']} ta\n"
        f"1 ta hujjat = {_price()} so'm\n\nSavol-javob va hujjat tekshirish — bepul.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(
            text="➕ Balansni to'ldirish", callback_data="pay")]]),
    )


@router.message(F.text == "📝 Hujjat tayyorlatish")
async def want_doc(message: Message, state: FSMContext) -> None:
    await state.set_state(St.doc_desc)
    await message.answer("Qanday hujjat kerak? Kimga va nima haqida ekanini yozing. Masalan:\n"
                         "«Ish beruvchiga 3 oylik ish haqi to'lanmagani bo'yicha talabnoma»\n"
                         "«Sudga aliment undirish haqida da'vo arizasi»")


@router.message(St.doc_desc, F.text)
async def doc_desc(message: Message, state: FSMContext) -> None:
    await state.clear()
    description = message.text[:300]
    await db.ensure_user(message.from_user.id)
    await db.set_pending_doc(message.from_user.id, description)
    await message.answer(f"«{description}»\n\nHujjat Word (.docx) faylda, qonun moddalariga havolalar bilan "
                         "tayyorlanadi. Bo'sh qolgan joylarni ([F.I.Sh.], [manzil]) o'zingiz to'ldirasiz.",
                         reply_markup=offer_kb())


# ---------------- paid documents ----------------

def payment_text() -> str:
    return (f"💳 To'lov: {_price()} so'm = 1 ta hujjat\n\n"
            f"Karta: {config.CARD_NUMBER}\n"
            f"Qabul qiluvchi: {config.CARD_HOLDER}\n\n"
            "1. Click, Payme, Uzum yoki bank ilovasi orqali shu kartaga pul o'tkazing.\n"
            "2. To'lov chekining skrinshotini shu yerga yuboring (summa, sana va qabul qiluvchi ko'rinsin).\n"
            "3. Chek avtomatik tekshiriladi va balansingiz to'ldiriladi.")


@router.callback_query(F.data == "pay")
async def pay(cb: CallbackQuery, state: FSMContext) -> None:
    await state.set_state(St.receipt)
    await cb.message.answer(payment_text())
    await cb.answer()


@router.callback_query(F.data == "draft")
async def draft(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    user = await db.ensure_user(cb.from_user.id)
    await cb.answer()
    if not user["pending_doc"]:
        await cb.message.answer("Avval qanday hujjat kerakligini yozing: «📝 Hujjat tayyorlatish».")
        return
    if not await db.spend_credit(cb.from_user.id):
        await state.set_state(St.receipt)
        await cb.message.answer("Hujjat tayyorlash pullik xizmat.\n\n" + payment_text())
        return
    await cb.message.answer("📝 Hujjat tayyorlanmoqda, 1 daqiqagacha vaqt olishi mumkin...")
    try:
        async with ChatActionSender.upload_document(bot=bot, chat_id=cb.message.chat.id):
            title, body = await rag.draft_document(cb.from_user.id, user["pending_doc"])
    except Exception as exc:
        await db.add_credits(cb.from_user.id, 1)  # refund: the user paid for a document they did not get
        await ai_error(cb.message, exc)
        return
    await db.set_pending_doc(cb.from_user.id, None)
    await cb.message.answer_document(
        BufferedInputFile(make_docx(title, body), filename=f"{safe_filename(title)}.docx"),
        caption="✅ Hujjat tayyor. [Kvadrat qavs] ichidagi joylarni to'ldiring va imzolang.\n\n" + DISCLAIMER,
    )


def safe_filename(title: str) -> str:
    return "".join(c if c.isalnum() or c in " -_" else "_" for c in title)[:60].strip() or "hujjat"


def make_docx(title: str, body: str) -> bytes:
    doc = Docx()
    heading = doc.add_heading(title, level=1)
    heading.alignment = 1  # center
    for para in body.split("\n"):
        doc.add_paragraph(para)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- receipts ----------------

@router.message(St.receipt, F.photo | F.document)
async def receipt(message: Message, state: FSMContext, bot: Bot) -> None:
    if message.photo:
        file, mime = message.photo[-1], "image/jpeg"
    elif message.document.mime_type in {"image/jpeg", "image/png", "image/webp", "application/pdf"}:
        file, mime = message.document, message.document.mime_type
    else:
        await message.answer("Chekni rasm (skrinshot) yoki PDF ko'rinishida yuboring.")
        return
    if (file.file_size or 0) > MAX_FILE:
        await message.answer("Fayl juda katta (20 MB dan oshmasin).")
        return
    if not await quota_ok(message):  # every receipt check is an AI call
        return

    data = (await bot.download(file)).read()
    image_hash = hashlib.sha256(data).hexdigest()
    try:
        async with ChatActionSender.typing(bot=bot, chat_id=message.chat.id):
            result = await payments.check_receipt(data, mime)
    except Exception as exc:
        await ai_error(message, exc)
        return

    txn_id = (result.get("transaction_id") or "").strip() or None
    if await db.receipt_seen(txn_id, image_hash):
        await message.answer("❌ Bu chek avval ishlatilgan.")
        return
    approved, reason, credits = payments.decide_now(result)
    status = "approved" if approved else "rejected"
    payment_id = await db.record_payment(message.from_user.id, int(result.get("amount") or 0) or None, credits,
                                         txn_id, image_hash, status, json.dumps(result, ensure_ascii=False))
    if approved and payment_id is None:
        await message.answer("❌ Bu chek avval ishlatilgan.")
        return
    if not approved:
        await message.answer(f"❌ {reason}")
        return

    await state.clear()
    total = await db.add_credits(message.from_user.id, credits)
    user = await db.ensure_user(message.from_user.id)
    kb = offer_kb() if user["pending_doc"] else None
    await message.answer(f"✅ {reason}! +{credits} ta kredit. Balans: {total} ta.", reply_markup=kb)
    await notify_admins(bot, data, mime, message, result, payment_id, credits)


async def notify_admins(bot: Bot, data: bytes, mime: str, message: Message, result: dict,
                        payment_id: int, credits: int) -> None:
    caption = (f"💰 To'lov #{payment_id}: {result.get('amount')} so'm → +{credits} kredit\n"
               f"User: {message.from_user.full_name} (id {message.from_user.id})\n"
               f"Qabul qiluvchi: {result.get('recipient_name')} *{result.get('recipient_card_last4')}\n"
               f"Vaqt: {result.get('paid_at')}  Tranzaksiya: {result.get('transaction_id')}")
    kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(
        text="❌ Soxta — bekor qilish", callback_data=f"revoke:{payment_id}")]])
    for admin in config.ADMIN_IDS:
        try:
            if mime.startswith("image/"):
                await bot.send_photo(admin, BufferedInputFile(data, "receipt.jpg"), caption=caption, reply_markup=kb)
            else:
                await bot.send_document(admin, BufferedInputFile(data, "receipt.pdf"), caption=caption, reply_markup=kb)
        except Exception:
            log.exception("could not notify admin %s", admin)


@router.callback_query(F.data.startswith("revoke:"))
async def revoke(cb: CallbackQuery, bot: Bot) -> None:
    if cb.from_user.id not in config.ADMIN_IDS:
        await cb.answer("Ruxsat yo'q", show_alert=True)
        return
    row = await db.revoke_payment(int(cb.data.split(":", 1)[1]))
    if not row:
        await cb.answer("Allaqachon bekor qilingan", show_alert=True)
        return
    await cb.answer("Bekor qilindi")
    await cb.message.edit_reply_markup(reply_markup=None)
    await bot.send_message(row["user_id"], "❌ To'lovingiz tasdiqlanmadi va kreditlar bekor qilindi. "
                                           "Xato bo'lsa, admin bilan bog'laning.")


# ---------------- admin ----------------

@router.message(Command("stats"), F.from_user.id.in_(config.ADMIN_IDS))
async def stats(message: Message) -> None:
    s = await db.stats()
    await message.answer(f"👥 Foydalanuvchilar: {s['users']}\n📚 Qonunlar: {s['laws']} ({s['chunks']} bo'lak, "
                         f"{s['embedded']} tasi ma'no-qidiruvga tayyor)\n"
                         f"💰 To'lovlar: {s['paid']} ta, {s['revenue']:,} so'm\n".replace(",", " ")
                         + f"🔖 Versiya: {os.environ.get('RENDER_GIT_COMMIT', 'local')[:7]}")


@router.message(Command("credit"), F.from_user.id.in_(config.ADMIN_IDS))
async def credit(message: Message, command: CommandObject) -> None:
    try:
        user_id, n = map(int, (command.args or "").split())
    except ValueError:
        await message.answer("Foydalanish: /credit <user_id> <soni>  (manfiy son — ayirish)")
        return
    await message.answer(f"Balans: {await db.add_credits(user_id, n)}")


# ---------------- AI chat & files ----------------

@router.message(F.photo | F.document)
async def file_message(message: Message, bot: Bot) -> None:
    if message.photo:
        file, mime = message.photo[-1], "image/jpeg"
    else:
        file, mime = message.document, message.document.mime_type or ""
        if mime not in FILE_MIMES:
            await message.answer("Bu fayl turini o'qiy olmayman. Rasm, PDF yoki Word (.docx) yuboring.")
            return
    if (file.file_size or 0) > MAX_FILE:
        await message.answer("Fayl juda katta (20 MB dan oshmasin). Muhim sahifalarni rasm qilib yuboring.")
        return
    if not await quota_ok(message):
        return
    await db.ensure_user(message.from_user.id)
    await message.answer("📄 Hujjatni o'qiyapman, biroz kuting...")
    try:
        async with ChatActionSender.typing(bot=bot, chat_id=message.chat.id):
            data = (await bot.download(file)).read()
            if mime == DOCX_MIME:
                text = "\n".join(p.text for p in Docx(io.BytesIO(data)).paragraphs)
                parts = [llm.types.Part(text=text[:60000])]
            else:
                parts = [llm.file_part(data, mime)]
            reply = await rag.analyze_file(message.from_user.id, parts, message.caption or "")
    except Exception as exc:
        await ai_error(message, exc)
        return
    await reply_with_offer(message, reply)


@router.message(F.text)
async def chat(message: Message, bot: Bot) -> None:
    if message.text.startswith("/"):
        await message.answer("Noma'lum buyruq. /start ni bosing.")
        return
    if not await quota_ok(message):
        return
    await db.ensure_user(message.from_user.id)
    await message.answer("⏳ Qonunlarni ko'rib chiqyapman, 15-30 soniya...")
    try:
        async with ChatActionSender.typing(bot=bot, chat_id=message.chat.id):
            reply = await rag.answer(message.from_user.id, message.text[:4000])
    except Exception as exc:
        await ai_error(message, exc)
        return
    await reply_with_offer(message, reply)


MAX_VOICE_SECONDS = 600


@router.message(F.voice | F.audio | F.video_note)
async def voice(message: Message, bot: Bot) -> None:
    media = message.voice or message.audio or message.video_note
    if (media.duration or 0) > MAX_VOICE_SECONDS:
        await message.answer("Ovozli xabar 10 daqiqadan oshmasin. Qisqaroq qilib yuboring.")
        return
    if (media.file_size or 0) > MAX_FILE:
        await message.answer("Fayl juda katta (20 MB dan oshmasin).")
        return
    if not await quota_ok(message):
        return
    await db.ensure_user(message.from_user.id)
    await message.answer("🎙 Ovozli xabaringizni tinglayapman...")
    mime = "video/mp4" if message.video_note else (getattr(media, "mime_type", None) or "audio/ogg")
    try:
        async with ChatActionSender.typing(bot=bot, chat_id=message.chat.id):
            data = (await bot.download(media)).read()
            text = await rag.transcribe(data, mime)
            if not text:
                await message.answer("Ovozni tushuna olmadim 😔 Iltimos, aniqroq gapirib qayta yuboring yoki yozib yuboring.")
                return
            await send_long(message, f"🎙 Siz dedingiz:\n«{text}»")
            reply = await rag.answer(message.from_user.id, text[:4000])
    except Exception as exc:
        await ai_error(message, exc)
        return
    await reply_with_offer(message, reply)


@router.message()
async def other(message: Message) -> None:
    await message.answer("Matn, ovozli xabar, rasm yoki PDF yuboring 🙂", reply_markup=MENU)
